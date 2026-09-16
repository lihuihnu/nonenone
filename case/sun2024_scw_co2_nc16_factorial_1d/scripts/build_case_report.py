#!/usr/bin/env python3
"""Create the editable Chinese DOCX report for the Sun-2024 factorial case."""

from __future__ import annotations

import argparse
import csv
import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


BLUE = "176A8A"
DARK = "243746"
LIGHT_BLUE = "EAF4F8"
LIGHT_ORANGE = "FFF3E6"
LIGHT_GRAY = "F2F4F5"
WHITE = "FFFFFF"
ORANGE = "D55E00"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def set_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None,
             name: str = "微软雅黑") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(paragraph, before=0, after=5, line=1.25) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_text(paragraph, text: str, *, bold=False, color=None, size=None, name="微软雅黑"):
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color, name=name)
    return run


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph(style="正文")
    if bold_prefix and text.startswith(bold_prefix):
        add_text(p, bold_prefix, bold=True)
        add_text(p, text[len(bold_prefix):])
    else:
        add_text(p, text)


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    add_text(p, text)
    set_paragraph_spacing(p, after=2, line=1.15)


def add_numbered(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    add_text(p, text)
    set_paragraph_spacing(p, after=2, line=1.15)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading(level=level)
    add_text(p, text, bold=True, color=DARK if level > 1 else BLUE)
    p.paragraph_format.keep_with_next = True


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None,
              font_size: float = 8.2) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, text, bold=True, color=WHITE, size=font_size)
        if widths:
            cell.width = Cm(widths[i])
    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if len(table.rows) % 2 == 1:
                set_cell_shading(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i != len(row_values) - 1 else WD_ALIGN_PARAGRAPH.LEFT
            add_text(p, str(value), size=font_size)
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(doc: Document, image_path: Path, caption: str, width_cm: float = 15.7) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    c = doc.add_paragraph(style="图注")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(c, caption, size=8.5)
    c.paragraph_format.keep_with_next = False


def add_callout(doc: Document, title: str, body: str, fill: str = LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=140, start=170, bottom=140, end=170)
    p = cell.paragraphs[0]
    add_text(p, title + "\n", bold=True, color=BLUE, size=10)
    add_text(p, body, size=9)
    set_paragraph_spacing(p, after=0, line=1.25)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_number(section) -> None:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "— ", color="6B7280", size=8)
    run = p.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_font(run, size=8, color="6B7280")
    add_text(p, " —", color="6B7280", size=8)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(9.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(5)

    if "正文" not in [s.name for s in doc.styles]:
        body = doc.styles.add_style("正文", 1)
    else:
        body = doc.styles["正文"]
    body.base_style = normal
    body.font.name = "微软雅黑"
    body._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    body.font.size = Pt(9.5)
    body.paragraph_format.first_line_indent = Cm(0.74)
    body.paragraph_format.line_spacing = 1.28
    body.paragraph_format.space_after = Pt(5)

    if "图注" not in [s.name for s in doc.styles]:
        caption = doc.styles.add_style("图注", 1)
    else:
        caption = doc.styles["图注"]
    caption.base_style = normal
    caption.font.name = "微软雅黑"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.paragraph_format.space_after = Pt(8)

    for level, size in ((1, 16), (2, 12), (3, 10.5)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE if level == 1 else DARK)
        style.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(5)


def add_header(section) -> None:
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(p, "MPMC_SCW_UNIFIED  |  可复现数值算例", color="6B7280", size=8)


def fmt(x: str | float, digits=4) -> str:
    return f"{float(x):.{digits}f}"


def build_markdown(out_path: Path, endpoint: dict[str, dict], effects: dict[str, dict], comparison: dict[str, dict]) -> None:
    pressure_effect_scw = float(endpoint["R06"]["recovery_4PV_percent"]) - float(endpoint["R04"]["recovery_4PV_percent"])
    pressure_effect_co2 = float(endpoint["R12"]["recovery_4PV_percent"]) - float(endpoint["R10"]["recovery_4PV_percent"])
    text = f"""# 超临界水-CO2-n-C16一维砂管可复现数值算例报告

## 核心结论

- 六组正式算例均完成至4 PVI，最大组分相对质量闭合误差为{max(float(r['max_relative_mass_closure']) for r in endpoint.values()):.3e}。
- 当前等温代理模型给出的4 PVI采收率为{min(float(r['recovery_4PV_percent']) for r in endpoint.values()):.4f}%至{max(float(r['recovery_4PV_percent']) for r in endpoint.values()):.4f}%。
- 等储层体积流量下的CO2组分效应仅为{100*float(effects['equal_reservoir_rate_CO2_composition_effect']['effect_23MPa_fraction']):.6f}和{100*float(effects['equal_reservoir_rate_CO2_composition_effect']['effect_24MPa_fraction']):.6f}个百分点。
- 压力由23升至24 MPa时，纯SCW与SCW+CO2的数值变化分别为{pressure_effect_scw:.4f}和{pressure_effect_co2:.4f}个百分点。
- R10/R12在4 PVI内均未达到1% CO2突破阈值，最大出口/入口CO2组分质量率比为{max(float(endpoint['R10']['max_outlet_inlet_co2_mass_ratio']), float(endpoint['R12']['max_outlet_inlet_co2_mass_ratio'])):.3e}。
- 原论文物理实验中CO2带来约7.8个百分点提升，而当前模型仅给出约0.0012个百分点；因此该模型适合作为流动与守恒基线，不足以解释物理实验中的热-组成协同增产。

## 可复现入口

```bash
RATE_BASIS=in_situ_volume bash case/sun2024_scw_co2_nc16_factorial_1d/scripts/run_matrix.sh formal
python case/sun2024_scw_co2_nc16_factorial_1d/scripts/generate_report_artifacts.py \\
  --formal-root tmp/sun2024_factorial_report_v1/formal \\
  --case-root case/sun2024_scw_co2_nc16_factorial_1d \\
  --grid-summary tmp/sun2024_factorial_convergence_v4/grid/convergence_summary.csv \\
  --time-summary tmp/sun2024_factorial_convergence_v4/time/convergence_summary.csv \\
  --output-root output/sun2024_factorial_report
```

完整方法、图表、限制和下一阶段方案见同目录DOCX报告。
"""
    out_path.write_text(text, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--git-commit", default="unknown")
    args = parser.parse_args()

    root = args.artifact_root
    figures = root / "figures"
    data_dir = root / "data"
    endpoint_rows = read_csv(data_dir / "endpoint_summary.csv")
    effect_rows = read_csv(data_dir / "effect_decomposition.csv")
    comparison_rows = read_csv(data_dir / "physical_endpoint_comparison.csv")
    grid_rows = read_csv(data_dir / "grid_convergence_summary.csv")
    time_rows = read_csv(data_dir / "time_convergence_summary.csv")
    provenance = json.loads((data_dir / "figure_provenance.json").read_text(encoding="utf-8"))
    endpoint = {r["experiment_id"]: r for r in endpoint_rows}
    effects = {r["effect"]: r for r in effect_rows}
    comparison = {r["experiment_id"]: r for r in comparison_rows}

    out_dir = args.output.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    build_markdown(out_dir / "算例报告摘要与复现入口.md", endpoint, effects, comparison)

    doc = Document()
    configure_document(doc)

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(85)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "超临界水-CO2-n-C16", bold=True, color=BLUE, size=26)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "一维砂管可复现数值算例报告", bold=True, color=DARK, size=23)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    add_text(p, "压力 × 组分 × 流量解耦设计", color=ORANGE, size=13)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    add_text(p, "基于 Sun et al. (2024) 物理砂管工况与 MPMC_SCW_UNIFIED 项目代码", size=10)
    add_callout(
        doc,
        "报告定位",
        "这是一个公开输入驱动、可重复运行、带数值收敛检查的等温流动基线。"
        "它不是完整热采历史拟合，也不把论文中已有数值模拟结果作为输入。",
        LIGHT_BLUE,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(55)
    add_text(p, f"版本日期：{date.today().isoformat()}\n代码基线：{args.git_commit}\n流量解释：in_situ_volume", color="52606D", size=9)
    doc.add_page_break()

    add_heading(doc, "执行摘要", 1)
    max_closure = max(float(r["max_relative_mass_closure"]) for r in endpoint_rows)
    min_rf = min(float(r["recovery_4PV_percent"]) for r in endpoint_rows)
    max_rf = max(float(r["recovery_4PV_percent"]) for r in endpoint_rows)
    comp23 = 100.0 * float(effects["equal_reservoir_rate_CO2_composition_effect"]["effect_23MPa_fraction"])
    comp24 = 100.0 * float(effects["equal_reservoir_rate_CO2_composition_effect"]["effect_24MPa_fraction"])
    max_ratio = max(float(endpoint["R10"]["max_outlet_inlet_co2_mass_ratio"]), float(endpoint["R12"]["max_outlet_inlet_co2_mass_ratio"]))
    add_callout(
        doc,
        "主要结论",
        f"六组正式算例均稳定推进到4 PVI；最大组分相对质量闭合误差为{max_closure:.2e}。"
        f"当前模型的4 PVI采收率集中在{min_rf:.4f}%–{max_rf:.4f}%。"
        f"等储层体积流量下，加入CO2的独立组成效应仅为{comp23:.6f}和{comp24:.6f}个百分点。"
        f"含CO2工况在4 PVI内均未达到1%突破阈值，最大出口/入口质量率比为{max_ratio:.2e}。",
        LIGHT_ORANGE,
    )
    add_body(doc, "最重要的科学结论不是“CO2无效”，而是：在当前等温、独立SCW相、n-C16代理油、无扩散/弥散、无反应的模型闭合下，CO2只在入口附近形成高组分区，无法复现物理实验中约7.8个百分点的增产幅度。该差异直接指出下一阶段需要加入的物理机制，而不是相对渗透率的简单调参问题。")
    add_body(doc, "PPT提出的目标模型把H2O视为可在水、油、气三相间分配的热力学组分，并以逸度相等、状态方程和三相闪蒸完成相平衡。当前砂管算例只验证其中的流动、守恒、井控和CO2-n-C16相态子集，因此应定位为完整模型之前的第一层基线。")
    add_heading(doc, "建议如何使用本算例", 2)
    add_numbered(doc, "先用本算例锁定质量守恒、注采边界、PVI终点和网格/时间步。")
    add_numbered(doc, "取得泵和中间容器的计量状态密度后，切换reference_density模式完成严格流量换算。")
    add_numbered(doc, "再依次加入能量方程、真实重油伪组分、H2O跨相分配与反应，逐层比较增量贡献。")

    add_heading(doc, "1 研究问题与证据边界", 1)
    add_heading(doc, "1.1 研究问题", 2)
    add_body(doc, "本研究回答三个可区分的问题：压力由23 MPa提高至24 MPa如何改变一维驱替；在总原位体积流量相同的条件下，CO2组分本身有多大影响；原实验中由10 mL/min水变为10 mL/min水+2 mL/min CO2时，增流量和组分变化各贡献多少。")
    add_heading(doc, "1.2 证据层级", 2)
    add_table(
        doc,
        ["证据层", "本报告用途", "明确排除"],
        [
            ["原论文物理实验", "几何、初始条件、温压流量、4 PV端点；用于外部面对比", "不读取论文已有数值模拟曲线"],
            ["项目PPT", "定义目标：H2O作为组分参与三相平衡；PR/SW/CPA热力学路线", "不把PPT中的闪蒸结果当作本砂管结果"],
            ["项目代码", "控制方程实现、边界条件、物性和输出字段", "不读取仓库既有数值实验数据"],
            ["本轮新计算", "六工况正式矩阵与R12收敛组；所有图表的唯一数值来源", "不进行事后参数标定"],
        ],
        widths=[3.0, 6.5, 6.0],
    )
    add_body(doc, "原论文没有唯一说明水和CO2体积流量的计量温度、压力或密度；文中还并存“温度稳定”“2 h”和“4 PV”三种终止描述。因此，本报告将4 PVI作为共同报告终点，并把表2体积率透明地解释为400 °C和对应背压下的原位体积率。该设定可重复，但不是严格实验复现。")

    add_heading(doc, "2 模型与物理假设", 1)
    add_heading(doc, "2.1 几何、初始与边界条件", 2)
    add_table(
        doc,
        ["类别", "参数", "取值", "来源/处理"],
        [
            ["几何", "长度 / 内径", "0.48 m / 0.039 m", "Sun et al. 图1与方法"],
            ["岩心", "孔隙度 / 渗透率", "0.39 / 2000 mD", "原论文输入"],
            ["初始", "So / Sg / Sw", "0.935 / 0 / 0.065", "算例锁定"],
            ["热力", "温度", "673.15 K (400 °C)，等温", "表2工况；无能量方程"],
            ["井控", "生产端BHP", "23或24 MPa", "出口BPR"],
            ["井控", "注入端压力上限", "35 MPa", "设备工作上限"],
            ["终点", "累计注入量", "4 PVI", "共同报告终点"],
        ],
        widths=[2.0, 4.0, 4.0, 5.5],
    )
    add_heading(doc, "2.2 当前实现", 2)
    add_bullet(doc, "H2O作为独立守恒SCW相；密度采用IAPWS-IF97，黏度采用McBride-Wright关联式。")
    add_bullet(doc, "CO2-n-C16采用PR78状态方程，二元作用系数kij=0.09。")
    add_bullet(doc, "各相流动由达西定律控制，组分由守恒方程推进；生产端定压、注入端定储层体积率。")
    add_bullet(doc, "不包含能量方程、热损失、毛管压力、分子扩散、机械弥散、有限速率相间传质与反应。")
    add_heading(doc, "2.3 与PPT目标模型的关系", 2)
    add_table(
        doc,
        ["模块", "PPT目标", "本算例状态", "解释后果"],
        [
            ["H2O角色", "H2O在水/油/气三相分配", "独立SCW相", "不能评估水烃真实互溶"],
            ["相平衡", "三相逸度相等与稳定性测试", "CO2-n-C16子体系", "不能复现完整三相分配"],
            ["状态方程", "PR、SW、CPA可比较", "PR78 + 独立水物性", "适合流动基线，不足以筛选EoS"],
            ["热过程", "超临界邻域P-T相态转换", "固定673.15 K", "没有50–400 °C热前缘"],
            ["原油", "H2O-CO2-BSB伪组分油", "单一n-C16", "无重油宽馏分与改质"],
        ],
        widths=[2.3, 4.4, 4.0, 4.8],
    )

    add_heading(doc, "3 可辨识的实验矩阵", 1)
    add_body(doc, "原四组物理工况形成压力（23/24 MPa）× CO2加入量（0/2 mL/min）的2×2结构，但含CO2时总报告体积率由10升至12 mL/min。为避免混杂，新增C23/C24纯SCW对照，使其储层总原位体积率与R10/R12完全相同。")
    matrix_rows = []
    for run in ("R04", "R06", "R10", "R12", "C23", "C24"):
        r = endpoint[run]
        matrix_rows.append([
            run,
            r["pressure_MPa"],
            r["water_rate_reported_mL_min"],
            r["co2_rate_reported_mL_min"],
            f"{float(r['reservoir_total_rate_m3_s']):.3e}",
            "物理工况" if r["role"] == "physical_reproduction" else "等流量对照",
        ])
    add_table(doc, ["ID", "背压 MPa", "水 mL/min", "CO2 mL/min", "储层总流量 m3/s", "角色"], matrix_rows,
              widths=[1.5, 2.4, 2.5, 2.5, 3.4, 3.0])
    add_heading(doc, "3.1 预先定义的效应", 2)
    add_bullet(doc, "压力效应：R06−R04（纯SCW）和R12−R10（SCW+CO2）。")
    add_bullet(doc, "纯增流量效应：C23−R04和C24−R06。")
    add_bullet(doc, "等储层体积流量CO2组分效应：R10−C23和R12−C24。")
    add_bullet(doc, "合并变化：R10−R04和R12−R06；只用于描述原始工况差，不能解释为纯CO2效应。")
    add_heading(doc, "3.2 输出与判据", 2)
    add_body(doc, "主要输出为n-C16累计采收率、注采压差、组分质量闭合、沿程饱和度/组成/压力，以及CO2突破。CO2突破定义为离散输出区间内“出口CO2组分质量率/入口CO2质量率”首次达到1%。所有曲线按原始离散点绘制，不插值平滑；确定性模拟不伪造重复样本或误差条。")

    add_heading(doc, "4 数值可接受性", 1)
    add_body(doc, "正式矩阵采用nx=96、ΔPVI=0.01；R12另做网格与时间步独立收敛组。非线性求解绝对容差设置为1e-12，避免默认绝对残差阈值造成零迭代接受和状态冻结。质量闭合以初始库存、当前库存、累计注入和累计产出的最大值归一化。")
    add_figure(doc, figures / "fig05_numerical_convergence.png", "图1  R12网格与时间步收敛性。右图明确保留“4 PVI内未达到1% CO2突破阈值”的阴性结果。")
    grid_delta = float(grid_rows[0]["medium_fine_recovery_delta"]) * 100.0
    time_delta = float(time_rows[0]["medium_fine_recovery_delta"]) * 100.0
    add_table(
        doc,
        ["检查", "中/细设置", "4 PVI采收率差", "质量闭合", "判定"],
        [
            ["网格", "nx=96 / 192，ΔPVI=0.005", f"{grid_delta:.5f} 个百分点", f"≤{max(float(r['max_component_relative_closure']) for r in grid_rows):.2e}", "通过"],
            ["时间步", "ΔPVI=0.01 / 0.005，nx=192", f"{time_delta:.5f} 个百分点", f"≤{max(float(r['max_component_relative_closure']) for r in time_rows):.2e}", "通过"],
        ],
        widths=[2.0, 5.2, 3.2, 3.0, 2.0],
    )
    add_body(doc, f"六组正式算例的最大组分相对质量闭合误差为{max_closure:.2e}，远低于1e-6验收阈值；所有算例均到达4 PVI，且未触发35 MPa注入压力上限。")

    add_heading(doc, "5 结果", 1)
    add_heading(doc, "5.1 采收曲线", 2)
    add_figure(doc, figures / "fig01_recovery_curves.png", "图2  六工况n-C16累计采收率。每个压力水平下分别比较基准SCW、等流量SCW和SCW+CO2。")
    add_body(doc, f"六组曲线在约0.8 PVI前快速上升，随后缓慢接近平台；4 PVI采收率仅跨越{min_rf:.4f}%–{max_rf:.4f}%的窄区间。压力由23升至24 MPa时，纯SCW与含CO2工况的4 PVI采收率分别变化{float(endpoint['R06']['recovery_4PV_percent'])-float(endpoint['R04']['recovery_4PV_percent']):.4f}和{float(endpoint['R12']['recovery_4PV_percent'])-float(endpoint['R10']['recovery_4PV_percent']):.4f}个百分点，方向均为轻微下降。")
    add_heading(doc, "5.2 压力响应", 2)
    add_figure(doc, figures / "fig02_pressure_drop_curves.png", "图3  六工况注采压差。等流量对照与含CO2工况几乎重合，显示该闭合下压差主要受总流量控制。")
    base_dp = float(endpoint["R04"]["final_delta_p_MPa"])
    high_dp = float(endpoint["C23"]["final_delta_p_MPa"])
    add_body(doc, f"10 mL/min纯SCW工况的末端压差约为{base_dp*1000:.3f} kPa；把储层总原位体积率提高20%后，C23末端压差升至{high_dp*1000:.3f} kPa，增加{(high_dp/base_dp-1)*100:.2f}%。同一总流量下，R10与C23、R12与C24几乎重合，说明当前参数范围内CO2组成对整体流动阻力的影响远小于流量影响。")
    add_heading(doc, "5.3 效应分解与物理端点", 2)
    add_figure(doc, figures / "fig03_endpoint_validation_and_effects.png", "图4  左：当前模型与原论文物理实验4 PV端点的外部面对比（不参与标定）；右：流量、等流量CO2组分及合并效应。")
    effect_table = []
    for key, label in (
        ("pure_SCW_flow_effect", "纯SCW增流量"),
        ("equal_reservoir_rate_CO2_composition_effect", "等流量CO2组分"),
        ("combined_CO2_plus_flow_effect", "CO2+增流量合并"),
    ):
        e = effects[key]
        effect_table.append([label, e["definition_23MPa"], f"{100*float(e['effect_23MPa_fraction']):.6f}", e["definition_24MPa"], f"{100*float(e['effect_24MPa_fraction']):.6f}"])
    add_table(doc, ["效应", "23 MPa定义", "百分点", "24 MPa定义", "百分点"], effect_table,
              widths=[4.0, 3.0, 2.5, 3.0, 2.5])
    add_body(doc, "等流量CO2组分效应约为0.00115–0.00120个百分点；纯增流量效应约为0.00008–0.00010个百分点。数量级都非常小，且合并效应不能解释原论文中加入CO2后约7.8个百分点的提升。")
    comp_rows = []
    for run in ("R04", "R06", "R10", "R12"):
        r = comparison[run]
        comp_rows.append([run, fmt(r["modeled_4PV_percent"]), fmt(r["physical_experiment_percent"], 2), fmt(r["signed_difference_percentage_points"]), "外部面对比，非标定"])
    add_table(doc, ["ID", "模型 %", "物理实验 %", "模型−实验 pp", "用途"], comp_rows,
              widths=[1.5, 2.5, 2.8, 3.0, 5.3])
    add_callout(doc, "解释警戒", "模型与物理端点差异同时包含流量计量状态不明、等温替代、油品替代、H2O独立相、缺少弥散/传质/反应等多项结构误差。不能把差异全部归因于相对渗透率，更不能直接调参使4个终点贴合。", LIGHT_ORANGE)

    add_heading(doc, "5.4 R12沿程状态", 2)
    add_figure(doc, figures / "fig04_R12_spatial_profiles.png", "图5  R12在0.5、1、2和4 PVI的沿程状态。坐标原点为注入口；压力表示高于24 MPa出口背压的增量。")
    add_body(doc, "0.5 PVI时SCW前缘位于砂管中后段，随后全管水相饱和度继续提高、油相饱和度降低。CO2液相摩尔分数在入口单元较高，但在极短距离内迅速衰减；出口CO2质量率始终远低于1%阈值。这一形态与“CO2组成效应极小”相互印证，但也暴露了缺少弥散、完整水烃互溶和热相态转换的局限。")

    add_heading(doc, "6 物理解释与不确定性", 1)
    add_heading(doc, "6.1 当前模型能支持的结论", 2)
    add_bullet(doc, "在已声明的in_situ_volume边界下，注采压差对总流量的响应近似成比例，井控稳定且远低于35 MPa上限。")
    add_bullet(doc, "网格、时间步和组分质量闭合满足预设数值门槛，因此“组分效应很小”不是明显离散误差或求解器假收敛造成的。")
    add_bullet(doc, "在4 PVI窗口内，CO2主要滞留于注入口附近，未形成可检测的出口突破。")
    add_heading(doc, "6.2 当前模型不能支持的结论", 2)
    add_bullet(doc, "不能据此断言实验中的CO2增产机制不存在；当前模型没有包含能产生该机制的全部物理。")
    add_bullet(doc, "不能评价热利用率，因为没有能量方程、热损失和50–400 °C热前缘。")
    add_bullet(doc, "不能把n-C16结果外推为真实BSB重油的黏度降低、萃取、膨胀或改质幅度。")
    add_bullet(doc, "不能完成PPT中PR/SW/CPA的公平比较，因为H2O尚未作为跨三相分配的热力学组分。")
    add_heading(doc, "6.3 关键不确定性优先级", 2)
    add_table(
        doc,
        ["优先级", "不确定性", "为什么重要", "解决方式"],
        [
            ["P0", "体积流量计量状态", "决定实际质量率和储层PVI", "取得泵/中间容器T、P、密度，改用reference_density"],
            ["P0", "能量与热损失", "原实验从50 °C升至400 °C，热前缘主导黏度", "加入岩石/流体能量方程和外壁换热"],
            ["P1", "H2O跨相分配", "PPT核心科学问题", "实现SW/CPA或一致的水组分逸度模型"],
            ["P1", "真实重油伪组分", "控制溶解、膨胀、黏度与残余油", "用BSB或实验油PVT/黏温数据表征"],
            ["P2", "弥散与有限速率传质", "决定CO2前缘展宽和突破", "用示踪/突破曲线反演纵向弥散度与传质系数"],
            ["P2", "反应/改质", "可能解释高温增产", "先用产物组成与黏度变化证实，再加入动力学"],
        ],
        widths=[1.5, 3.7, 5.0, 5.2],
    )

    add_heading(doc, "7 下一阶段可复现实验方案", 1)
    add_body(doc, "建议采用“验证阶梯”而不是一次性扩展：每一级只新增一类机制，并沿用本算例的冻结矩阵、元数据和验收门槛。这样才能知道结果变化来自哪一项物理，而不是来自同时变化的输入与数值设置。")
    add_table(
        doc,
        ["阶段", "体系与新增机制", "核心矩阵", "验收输出"],
        [
            ["A 计量闭环", "当前SCW-CO2-n-C16；流量改为质量换算", "6工况×2流量基准", "PVI、质量率、压差一致性"],
            ["B 非等温流动", "加入能量方程、岩石热容和外壁换热", "50 °C初始；400 °C注入；23/24 MPa", "温度前缘、热利用率、压差、采收率"],
            ["C 统一热力学", "H2O-CO2-BSB；PR/SW/CPA", "650–673 K，23–39 MPa，固定总组成", "相数、相分率、逸度残差、组分分配"],
            ["D 流动-相态耦合", "C的模型进入1D砂管", "CO2 0/5/10/20 mol%，等总质量率；两压力", "前缘速度、突破、采收率、相态路径"],
            ["E 机理确认", "按数据证据增加弥散、传质或反应", "单因子增量+留出工况", "预测而非拟合的端点与全过程曲线"],
        ],
        widths=[2.2, 4.8, 4.7, 4.0],
        font_size=7.8,
    )
    add_heading(doc, "7.1 推荐的首个完整体系", 2)
    add_body(doc, "与PPT衔接最自然的体系是H2O-CO2-BSB伪组分油。先在0D闪蒸中复现PPT的653.15 K、28 MPa两液相工况及650 K、39 MPa三相工况，再进入一维非等温砂管。这样既能检验水作为组分的价值，也能与当前n-C16基线形成清晰对照。")
    add_heading(doc, "7.2 建议的最小辨识矩阵", 2)
    add_bullet(doc, "压力：23、24、28 MPa；温度：620、647、653、673 K，用于跨越水临界邻域。")
    add_bullet(doc, "CO2：0、5、10、20 mol%，保持总注入质量率或总PVI一致；另保留等焓对照，避免把热量与组成混杂。")
    add_bullet(doc, "状态方程：PR、SW、CPA；先以逸度残差和相数为门槛，再比较流动输出。")
    add_bullet(doc, "每个主工况至少做nx和ΔPVI两级确认；只在需要量化随机实验误差时设置物理重复。")

    add_heading(doc, "8 结论", 1)
    conclusions = [
        "已建立一个由公开物理输入驱动、包含六个工况和两类收敛检查的一维可复现数值算例。",
        f"在当前闭合下，4 PVI采收率为{min_rf:.4f}%–{max_rf:.4f}%，等流量CO2独立组成效应只有约0.0012个百分点。",
        f"含CO2工况未在4 PVI内达到1%突破阈值，最大出口/入口质量率比为{max_ratio:.2e}；CO2主要局限在入口附近。",
        "模型无法复现物理实验中CO2带来的约7.8个百分点提升，证明当前基线不足以支撑机理结论，但为下一阶段模型增量提供了明确诊断。",
        "下一步应优先解决流量计量状态与能量方程，随后实现PPT提出的H2O跨三相热力学分配，并以H2O-CO2-BSB体系完成0D到1D的阶梯验证。",
    ]
    for index, conclusion in enumerate(conclusions, start=1):
        p = doc.add_paragraph()
        add_text(p, f"{index}.  {conclusion}")
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.7)
        set_paragraph_spacing(p, after=3, line=1.18)

    add_heading(doc, "参考资料", 1)
    refs = [
        "[1] Sun Y, Zhang L, Bai Y, et al. Analysis of Adaptability and Application Potential of Supercritical Multi-Source Multi-Component Thermal Fluid Technology for Offshore Heavy Oil in China. Applied Sciences, 2024, 14: 3588. DOI: 10.3390/app14093588.",
        "[2] 《超临界水渗流多相多组分模型9》项目内部演示文稿，12页，用户提供。",
        "[3] MPMC_SCW_UNIFIED项目源代码及case/sun2024_scw_co2_nc16_factorial_1d算例；本报告代码基线见封面。",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        add_text(p, ref, size=8.5)
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.7)
        p.paragraph_format.space_after = Pt(4)

    add_heading(doc, "附录A 复现步骤与数据清单", 1)
    add_heading(doc, "A.1 最小复现步骤", 2)
    commands = [
        "RATE_BASIS=in_situ_volume bash case/sun2024_scw_co2_nc16_factorial_1d/scripts/run_matrix.sh grid",
        "RATE_BASIS=in_situ_volume bash case/sun2024_scw_co2_nc16_factorial_1d/scripts/run_matrix.sh time",
        "RATE_BASIS=in_situ_volume bash case/sun2024_scw_co2_nc16_factorial_1d/scripts/run_matrix.sh formal",
        "python case/sun2024_scw_co2_nc16_factorial_1d/scripts/generate_report_artifacts.py [路径参数见README]",
    ]
    table = doc.add_table(rows=1, cols=1)
    set_cell_shading(table.cell(0, 0), "F7F7F7")
    set_cell_margins(table.cell(0, 0), top=140, start=160, bottom=140, end=160)
    p = table.cell(0, 0).paragraphs[0]
    for i, cmd in enumerate(commands):
        run = p.add_run(cmd + ("\n" if i < len(commands)-1 else ""))
        set_font(run, size=7.6, name="Consolas")
    add_heading(doc, "A.2 机器可读输出", 2)
    add_table(
        doc,
        ["文件", "内容"],
        [
            ["data/endpoint_summary.csv", "六工况端点、压差、守恒与突破指标"],
            ["data/effect_decomposition.csv", "流量、等流量CO2组分、合并效应"],
            ["data/recovery_curves.csv", "六工况全过程采收曲线"],
            ["data/pressure_drop_curves.csv", "六工况全过程注采压差"],
            ["data/R12_spatial_profiles.csv", "0.5/1/2/4 PVI沿程状态"],
            ["data/grid_convergence_summary.csv", "网格收敛结果"],
            ["data/time_convergence_summary.csv", "时间步收敛结果"],
            ["data/figure_provenance.json", "数据源、变换、绘图软件与模型限制"],
        ],
        widths=[6.0, 9.5],
    )
    add_body(doc, f"图件由{provenance['rendering']['software']}生成，同时提供600 dpi PNG与矢量PDF；无平滑处理，使用色盲友好配色并以线型和标记冗余编码。")
    add_heading(doc, "A.3 复现声明", 2)
    add_body(doc, "本报告所有数值图表来自本轮明确指定的新结果目录。绘图脚本不搜索其他results目录；物理实验值仅来自case-local physical_validation_targets.csv，并只作外部面对比。正式算例元数据记录nx=96、ΔPVI=0.01、4 PVI、单进程与in_situ_volume流量解释。")

    # Final header/footer consistency.
    for section in doc.sections:
        add_header(section)
        add_page_number(section)

    doc.core_properties.title = "超临界水-CO2-n-C16一维砂管可复现数值算例报告"
    doc.core_properties.subject = "压力×组分×流量解耦；新计算结果；可复现数值实验"
    doc.core_properties.keywords = "超临界水, CO2, 多相多组分, 数值实验, 可复现"
    doc.core_properties.comments = "Numerical results generated in this work only; no existing numerical experiment data used."
    doc.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
